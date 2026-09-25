"""Prepare repository accession lists from the supplied search exports."""

import argparse
import json
import re
import shutil
from pathlib import Path

import openpyxl

from metadata_io import read_csv, write_csv


def full_list(sheet, pattern):
    start = next(row[0].row for row in sheet if str(row[0].value).strip() == "FULL LIST")
    return [row for row in sheet.iter_rows(min_row=start + 1)
            if re.fullmatch(pattern, str(row[0].value))]


def prepare(workbook_path, dbgap_path, root):
    source_dir = root / "inputs" / "original"
    source_dir.mkdir(parents=True, exist_ok=True)
    for path in [workbook_path, dbgap_path]:
        target = source_dir / path.name
        if path.resolve() != target.resolve():
            shutil.copy2(path, target)
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    pride = [row[0].value for row in full_list(workbook["PRIDE"], r"(?:PXD|PRD)\d+")]
    massive = []
    for row in full_list(workbook["MassIVE"], r"MSV\d+"):
        accession, alias = row[0].value, row[1].value or ""
        massive.append({"dataset_id": accession, "source": "massive", "px_accession": alias,
                        "in_pride_list": alias in pride if alias else False,
                        "source_url": f"https://massive.ucsd.edu/ProteoSAFe/dataset.jsp?accession={accession}"})
    write_csv(root / "inputs/massive_ids.csv", massive)
    write_csv(root / "inputs/pride_reference_ids.csv", [{"dataset_id": value} for value in pride])

    sheet = workbook["dbgap"]
    headers = [cell.value for cell in sheet[2]]
    old_records = {}
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if re.fullmatch(r"phs\d+\.v\d+\.p\d+", str(row[0])) and row[1]:
            old_records[row[0]] = dict(zip(headers, row))
    old_ids = [row[0].value for row in full_list(sheet, r"phs\d+\.v\d+\.p\d+")]
    new_records = {row["accession"]: row for row in read_csv(dbgap_path)}
    # Match the study, not its version, so an updated accession is not a new study.
    old_by_base = {value.split(".")[0]: value for value in old_ids}
    new_by_base = {value.split(".")[0]: value for value in new_records}
    reconciled = []
    for base in sorted(old_by_base.keys() | new_by_base.keys()):
        previous, current = old_by_base.get(base, ""), new_by_base.get(base, "")
        accession = current or previous
        metadata = new_records.get(accession, old_records.get(accession, {}))
        reconciled.append({"dataset_id": accession, "source": "dbgap", "study_id": base,
                           "workbook_accession": previous, "csv_accession": current,
                           "in_workbook": bool(previous), "in_new_csv": bool(current),
                           "version_changed": bool(previous and current and previous != current),
                           "title": metadata.get("name", ""),
                           "review_status": "not_reviewed",
                           "source_url": f"https://www.ncbi.nlm.nih.gov/projects/gap/cgi-bin/study.cgi?study_id={accession}"})
    write_csv(root / "inputs/dbgap_ids.csv", reconciled)
    (root / "inputs/dbgap_search_metadata.json").write_text(
        json.dumps({**old_records, **new_records}, indent=2, default=str), encoding="utf-8")
    workbench = [{"dataset_id": row[0], "source": "metabolomics_workbench", "title": row[1],
                  "sample_source": row[2], "organism": row[3], "condition": row[4] or ""}
                 for row in workbook["MetabolomicsWorkbench"].iter_rows(min_row=2, values_only=True)]
    write_csv(root / "inputs/workbench_ids.csv", workbench)
    terms = [{"search_term": row[0], "reported_hits": row[1]}
             for row in workbook["MetaboLights"].iter_rows(values_only=True)]
    write_csv(root / "inputs/metabolights_search_terms.csv", terms)

    # This pilot deliberately includes different layouts and ambiguous search hits.
    selections = {
        "MSV000086385": "Placental tissue; imported dataset with an exact PRIDE cross-reference.",
        "MSV000085995": "Original trophoblast subtype study; compare sample and method detail.",
        "MSV000095456": "Mass spectrometry imaging; test mapping beyond conventional proteomics.",
        "MSV000094962": "Lipidomics with placenta and maternal blood; test tissue-specific counts.",
        "MSV000100811": "Title refers to germline RNA and placenta; test relevance and assay assumptions.",
        "phs001320.v1.p1": "Chorionic villi study with a public molecular-data table.",
        "phs001782.v2.p1": "Large cohort with multiple tissues; distinguish subjects and samples.",
        "phs001886.v6.p1": "Study version changed between the two supplied lists.",
        "phs003122.v1.p1": "Present only in the older list; assess possible loss from a narrower search.",
        "phs003002.v2.p1": "Pancreatic cancer search hit; investigate pregnancy-test terminology.",
    }
    by_id = {row["dataset_id"]: row for row in massive + reconciled}
    pilot = [{"dataset_id": key, "source": by_id[key]["source"],
              "source_url": by_id[key]["source_url"], "selection_reason": reason}
             for key, reason in selections.items()]
    write_csv(root / "inputs/pilot_ids.csv", pilot)
    inventory = {"massive": len(massive), "pride_reference": len(pride),
                 "massive_pride_matches": sum(row["in_pride_list"] for row in massive),
                 "dbgap_workbook": len(old_ids), "dbgap_new_csv": len(new_records),
                 "dbgap_reconciled_candidates": len(reconciled),
                 "dbgap_shared_studies": len(old_by_base.keys() & new_by_base.keys()),
                 "dbgap_version_changes": sum(row["version_changed"] for row in reconciled),
                 "metabolomics_workbench": len(workbench),
                 "metabolights_search_terms": len(terms), "metabolights_unique_studies": None,
                 "pilot_records": len(pilot)}
    (root / "inputs/inventory.json").write_text(json.dumps(inventory, indent=2) + "\n")
    print(json.dumps(inventory, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--dbgap", type=Path, required=True)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    prepare(args.workbook, args.dbgap, args.root)
